"""Exercise reusable template authoring, validation, application and native structure."""
from __future__ import annotations
import argparse
import hashlib
import json
from pathlib import Path
import shutil
import sys
import time
import zipfile
from xml.etree import ElementTree as ET

CAP = Path(__file__).resolve().parents[1]
ROOT = CAP.parents[1]
sys.path.insert(0, str(ROOT / 'scripts'))
from design_brain_process_utils import run_command

def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--work-dir', required=True)
    args=parser.parse_args()
    work=Path(args.work_dir).resolve()
    if work.exists(): raise SystemExit('Use a new work directory; fixtures are never overwritten.')
    work.mkdir(parents=True)
    checks=[]
    started=time.monotonic()
    def cli(label,*argv,expected=0):
        remaining=240-(time.monotonic()-started)
        if remaining<=0: raise AssertionError('Template smoke global timeout')
        result=run_command([sys.executable,'-B',str(CAP/'runtime/presentation.py'),*map(str,argv)],cwd=work,timeout=min(90,remaining),isolate_process_group=True)
        (work/(label.replace(' ','-')+'.log')).write_text(result.stdout+'\n'+result.stderr,encoding='utf-8')
        assert not result.timedOut and ((result.returncode==0)==(expected==0)), result.stderr+result.stdout[-9000:]
        checks.append(label);print('PASS '+label,flush=True)
        return json.loads(result.stdout or result.stderr)
    def init(name):
        project=work/(name+'_20261012')
        cli('initialize '+name,'init',project,'--mode','quick')
        (project/'templates').mkdir(exist_ok=True)
        return project
    brand=init('brand')
    (brand/'templates/design_spec.brand.test_identity.md').write_text('''---
brand_id: test_identity
kind: brand
summary: Neutral engineering fixture identity for reusable-template tests.
primary_color: "#174A67"
---
# Test Identity
## I. Brand Overview
An engineering fixture with a user-specified identity; no organization is represented.
## II. Color Scheme
| Role | Color | Provenance |
|---|---|---|
| Primary | #174A67 | user |
| Background | #FFFFFF | user |
| Text | #1F2937 | user |
## III. Typography
Arial for titles and body; explicit scale is resolved by the consuming deck.
## IV. Logo
No logo; use a text clue placeholder when a consumer supplies its brand.
## V. Voice & Tone
Direct and evidence-led.
## VI. Icon Style
One coherent outline family when icons aid understanding.
''',encoding='utf-8')
    cli('validate authored brand','template-check',brand)
    style=init('style')
    source_style=CAP/'resources/styles/consulting-decision/templates/design_spec.md'
    style_text=source_style.read_text(encoding='utf-8').replace('style_id: consulting-decision','style_id: test_method')
    (style/'templates/design_spec.style.test_method.md').write_text(style_text,encoding='utf-8')
    cli('validate adopted style','template-check',style)
    layout=init('layout')
    source_layout=CAP/'resources/layouts/presentation_core/templates'
    source_hashes={p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in source_layout.iterdir() if p.is_file()}
    for p in source_layout.glob('*.svg'): shutil.copy2(p,layout/'templates'/p.name)
    layout_spec=(source_layout/'design_spec.md').read_text(encoding='utf-8').replace('layout_id: presentation_core','layout_id: test_structure')
    (layout/'templates/design_spec.layout.test_structure.md').write_text(layout_spec,encoding='utf-8')
    cli('validate twenty role layouts','template-check',layout)
    preview=work/'template-preview.pptx'
    cli('export structured template preview','template-preview',layout,'--output',preview)
    cli('reject template preview overwrite','template-preview',layout,'--output',preview,expected=1)
    from pptx import Presentation
    deck=Presentation(str(preview))
    layouts=[l for m in deck.slide_masters for l in m.slide_layouts]
    assert len(deck.slides)==20 and len(layouts)==20 and len(deck.slide_masters)>=1
    assert len({l.name for l in layouts})==20
    assert sum(len(s.placeholders) for s in deck.slides)>=40
    for slide in deck.slides:
        expected={p.placeholder_format.idx:(p.left,p.top,p.width,p.height) for p in slide.slide_layout.placeholders}
        assert {p.placeholder_format.idx:(p.left,p.top,p.width,p.height) for p in slide.placeholders}==expected
    checks.append('twenty named native layouts and matching editable placeholder frames')
    # A content edit exercises real reusable native text frames; schema checks
    # alone would miss broken slide/layout relationships.
    target=next(s for s in deck.slides[0].placeholders if s.has_text_frame)
    target.text_frame.paragraphs[0].runs[0].text='A longer replacement title used to verify an editable native placeholder'
    edited=work/'template-content-edit.pptx';deck.save(str(edited))
    reopened=Presentation(str(edited))
    assert any('replacement title' in s.text for s in reopened.slides[0].shapes if s.has_text_frame)
    checks.append('native placeholder text replacement survives reopen')
    combined=init('combined')
    cli('apply brand style and layout','apply-template',combined,'--root',brand,'--root',style,'--root',layout)
    cli('validate combined reusable workspace','template-check',combined)
    assert len(list((combined/'templates').glob('design_spec.*.md')))==3
    applied=work/'applied-template.pptx'
    cli('preview applied reusable workspace','template-preview',combined,'--output',applied)
    assert len(Presentation(str(applied)).slides)==20
    # A Deck kind owns identity/application as well as structure. Author its
    # qualified spec over the same engineering fixture vocabulary, not a second
    # marketplace product or a claimed original design.
    full=init('deck')
    for p in (layout/'templates').glob('*.svg'): shutil.copy2(p,full/'templates'/p.name)
    deck_spec=layout_spec.replace('layout_id: test_structure','deck_id: test_deck').replace('kind: layout','kind: deck')
    deck_spec+='\n## I. Template Overview\nEngineering demonstration of editable status reporting.\n\n## II. Color Scheme\nPrimary #174A67; text #1F2937; background #FFFFFF.\n\n## III. Typography\nArial titles and body, maintaining the authored readable hierarchy.\n'
    (full/'templates/design_spec.deck.test_deck.md').write_text(deck_spec,encoding='utf-8')
    cli('validate authored deck kind','template-check',full)
    cli('preview deck kind','template-preview',full,'--output',work/'deck-template.pptx')
    assert source_hashes=={p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in source_layout.iterdir() if p.is_file()}
    checks.append('bundled reusable source remains unchanged')
    result={'status':'PASS','checks':checks,'slides':20,'layouts':20,'workDir':str(work),'visualInspection':'requires-render-review'}
    (work/'result.json').write_text(json.dumps(result,indent=2)+'\n',encoding='utf-8')
    print(json.dumps(result,indent=2))
    return 0

if __name__=='__main__': raise SystemExit(main())

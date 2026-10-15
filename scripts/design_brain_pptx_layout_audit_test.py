"""Focused layout-audit regressions. Fixtures are native, synthetic PPTX files.

Uses the already installed python-pptx for fixture authoring only. Run with the
bundled Python: python -B scripts/design_brain_pptx_layout_audit_test.py
"""
from __future__ import annotations

import base64
import copy
from io import BytesIO
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from zipfile import ZipFile

sys.dont_write_bytecode = True
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.util import Inches
from PIL import Image
from pptx.chart.data import ChartData
from pptx.enum.chart import XL_CHART_TYPE
import design_brain_pptx_layout_audit as audit

PNG = base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAusB9Wl6hVgAAAAASUVORK5CYII=")
BASE = [("text", (0.7, 0.5, 3.5, 0.7)), ("text", (0.7, 1.5, 3.5, 2.5)), ("picture", (5.0, 1.0, 4.0, 3.4))]


class LayoutAuditTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="pptx-layout-audit-")
        self.addCleanup(self.temp.cleanup)
        self.base = Path(self.temp.name).resolve()

    def deck(self, filename="current.pptx", slides=None, payload="Original", color="112233", form=MSO_SHAPE.RECTANGLE, scale=1):
        prs = Presentation()
        prs.slide_width, prs.slide_height = Inches(10 * scale), Inches(5.625 * scale)
        for regions in [BASE] if slides is None else slides:
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            for role, box in regions:
                geometry = [Inches(value * scale) for value in box]
                if role == "picture":
                    image = BytesIO(PNG)
                    if payload != "Original":
                        image = BytesIO()
                        Image.new("RGB", (2, 2), "red").save(image, format="PNG")
                        image.seek(0)
                    shape = slide.shapes.add_picture(image, *geometry)
                    shape.crop_left = 0.05 if payload != "Original" else 0
                elif role == "text":
                    shape = slide.shapes.add_textbox(*geometry)
                    shape.text = payload
                    shape.text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor.from_string(color)
                else:
                    shape = slide.shapes.add_shape(form, *geometry)
                    shape.fill.solid()
                    shape.fill.fore_color.rgb = RGBColor.from_string(color)
        path = self.base / filename
        prs.save(path)
        return path

    def fp(self, path, deck_id="deck"):
        return audit.extract(path, deck_id)

    def compare(self, current, previous):
        return audit.compare(self.fp(current, "new"), [self.fp(path, f"old-{index}") for index, path in enumerate(previous)])

    def cli(self, *args):
        result = subprocess.run([sys.executable, "-B", audit.__file__, *map(str, args)], capture_output=True, text=True, timeout=20)
        self.assertNotIn("Traceback", result.stderr)
        return result.returncode, json.loads(result.stdout)

    def test_text_color_photo_crop_swaps_and_shape_forms_do_not_make_new_layout(self):
        shapes = BASE + [("shape", (0.6, 4.7, 8.4, 0.4))]
        old = self.deck("old.pptx", [shapes])
        new = self.deck("new.pptx", [shapes], payload="Different copy and picture crop", color="FEDCBA", form=MSO_SHAPE.OVAL)
        result = self.compare(new, [old])
        self.assertEqual(result["status"], "REJECT_IDENTICAL")
        self.assertTrue(result["comparisons"][0]["identical_outer_composition"])
        self.assertEqual(result["design_readiness"], "UNVERIFIED")
        self.assertNotEqual(result["current"]["source"]["sha256"], result["previous"][0]["source"]["sha256"])

    def test_renamed_file_scaled_canvas_and_shape_id_changes_are_not_originality(self):
        old = self.deck("old.pptx")
        new = self.deck("renamed.pptx", scale=2)
        prs = Presentation(new)
        for shape in prs.slides[0].shapes:
            shape._element.xpath(".//p:cNvPr")[0].set("id", str(shape.shape_id + 50))
        prs.save(new)
        self.assertEqual(self.compare(new, [old])["status"], "REJECT_IDENTICAL")

    def test_all_slides_across_all_history_decks_are_compared(self):
        moved = [(role, (box[0] + 0.3, box[1] + 0.2, box[2], box[3])) for role, box in BASE]
        old = self.deck("old.pptx", [BASE, moved])
        other = self.deck("other.pptx", [moved])
        new = self.deck("new.pptx", [moved, BASE])
        result = self.compare(new, [old, other])
        self.assertEqual(result["compared_pair_count"], 6)
        clones = {(pair["current"]["slide_number"], pair["previous"]["slide_number"]) for pair in result["comparisons"] if pair["status"] == "REJECT_IDENTICAL" and pair["previous"]["deck_id"] == "old-0"}
        self.assertEqual(clones, {(1, 2), (2, 1)})

    def test_size_type_and_decoration_count_changes_still_need_main_review(self):
        old = self.deck("old.pptx", [BASE + [("shape", (0.6, 4.7, 1.0, 0.1))]])
        resized = [(role, (box[0], box[1], box[2] * 0.8, box[3] * 0.75)) for role, box in BASE]
        new = self.deck("new.pptx", [resized + [("shape", (8.9, 4.7, 0.15, 0.15)), ("shape", (9.2, 4.7, 0.15, 0.15))]], form=MSO_SHAPE.OVAL)
        result = self.compare(new, [old])
        self.assertEqual(result["status"], "MAIN_REVIEW_REQUIRED")
        self.assertTrue(result["comparisons"][0]["likely_repeated_composition"])
        self.assertEqual(result["comparisons"][0]["matched_anchor_count"], 3)

    def test_materially_moved_composition_is_not_false_identical_or_automatic_pass(self):
        old = self.deck("old.pptx")
        moved = [("text", (5.5, 0.5, 3.5, 0.7)), ("text", (5.5, 1.5, 3.5, 2.5)), ("picture", (0.5, 1.0, 4.0, 3.4))]
        result = self.compare(self.deck("new.pptx", [moved]), [old])
        self.assertEqual(result["status"], "MAIN_REVIEW_REQUIRED")
        self.assertFalse(result["comparisons"][0]["identical_outer_composition"])
        self.assertFalse(result["comparisons"][0]["likely_repeated_composition"])

    def test_table_chart_outer_frames_and_empty_shape_textbodies_are_not_full_layout_proof(self):
        prs = Presentation()
        prs.slide_width, prs.slide_height = Inches(10), Inches(5.625)
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(10), Inches(5.625))
        slide.shapes.add_table(2, 2, Inches(0.5), Inches(1), Inches(3), Inches(3))
        data = ChartData()
        data.categories = ["A", "B"]
        data.add_series("Fixture", [1, 2])
        slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(5), Inches(1), Inches(4), Inches(3), data)
        path = self.base / "opaque-frames.pptx"
        prs.save(path)
        fp = self.fp(path)
        self.assertEqual(fp["status"], "EXTRACTED")
        background = fp["slides"][0]["shapes"][0]
        self.assertEqual((background["role"], background["region_role"]), ("shape", "background"))
        self.assertEqual({shape["role"] for shape in fp["slides"][0]["shapes"]}, {"shape", "table", "chart"})
        pair = audit.compare(fp, [fp])["comparisons"][0]
        self.assertEqual(pair["status"], "MAIN_REVIEW_REQUIRED")
        self.assertTrue(pair["outer_frames_match"])
        self.assertFalse(pair["identical_outer_composition"])

    def test_disjoint_xml_z_order_and_group_ids_are_evidence_not_layout_identity(self):
        old = self.deck("old.pptx")
        new = self.deck("new.pptx")
        prs = Presentation(new)
        tree = prs.slides[0].shapes._spTree
        shapes = list(prs.slides[0].shapes)
        for shape in reversed(shapes):
            tree.remove(shape._element)
            tree.append(shape._element)
        prs.save(new)
        self.assertEqual(self.compare(new, [old])["status"], "REJECT_IDENTICAL")

    def test_native_group_scaling_is_resolved_but_rotated_group_blocks(self):
        prs = Presentation()
        prs.slide_width, prs.slide_height = Inches(10), Inches(5.625)
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        group = slide.shapes.add_group_shape()
        child = group.shapes.add_textbox(Inches(1), Inches(1), Inches(1), Inches(1))
        child.text = "Grouped text"
        group.left, group.top, group.width, group.height = Inches(2), Inches(1), Inches(2), Inches(2)
        path = self.base / "group.pptx"
        prs.save(path)
        fp = self.fp(path)
        self.assertEqual(fp["status"], "EXTRACTED")
        box = fp["slides"][0]["shapes"][0]["box"]
        self.assertEqual(box[:3], [0.2, 0.177778, 0.2])
        self.assertTrue(fp["slides"][0]["shapes"][0]["group_origin"])
        group.rotation = 90
        prs.save(path)
        fp = self.fp(path)
        self.assertEqual(fp["status"], "BLOCKED")
        self.assertIn("UNSUPPORTED_GROUP_ROTATION_OR_FLIP", fp["slides"][0]["issues"])

    def test_layout_placeholder_and_static_master_geometry_are_included(self):
        prs = Presentation()
        prs.slide_width, prs.slide_height = Inches(10), Inches(5.625)
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        slide.shapes.title.text = "Inherited title"
        title = slide.shapes.title
        expected = [round(float(value) / total, 6) for value, total in zip((title.left, title.top, title.width, title.height), (prs.slide_width, prs.slide_height, prs.slide_width, prs.slide_height))]
        shape = slide.shapes.add_textbox(Inches(0.3), Inches(5), Inches(2), Inches(0.3))
        shape.text = "Master footer"
        master_tree = prs.slide_master.shapes._spTree
        master_tree.append(copy.deepcopy(shape._element))
        slide.shapes._spTree.remove(shape._element)
        path = self.base / "inheritance.pptx"
        prs.save(path)
        fp = self.fp(path)
        self.assertEqual(fp["status"], "EXTRACTED")
        native_title = next(shape for shape in fp["slides"][0]["shapes"] if shape["origin"] == "slide" and shape["role_hint"] == "ctrTitle")
        self.assertEqual(native_title["box"], expected)
        footer = next(shape for shape in fp["slides"][0]["shapes"] if shape["origin"] == "master")
        self.assertEqual(footer["box"][:2], [0.03, 0.888889])

    def test_empty_flattened_and_unresolved_geometry_are_blocked(self):
        for name, slides in (("empty", [[]]), ("flat", [[("picture", (0, 0, 10, 5.625))]])):
            with self.subTest(name=name):
                fp = self.fp(self.deck(name + ".pptx", slides))
                self.assertEqual(fp["status"], "BLOCKED")
        path = self.deck("unknown.pptx")
        prs = Presentation(path)
        shape = prs.slides[0].shapes[0]
        shape._element.spPr.remove(shape._element.spPr.xfrm)
        prs.save(path)
        fp = self.fp(path)
        self.assertIn("UNRESOLVED_ZERO_GEOMETRY", fp["slides"][0]["issues"])

    def test_missing_history_never_proves_new_layout(self):
        result = audit.compare(self.fp(self.deck()), [])
        self.assertEqual((result["status"], result["design_readiness"]), ("NO_HISTORY", "UNVERIFIED"))
        self.assertEqual(result["compared_pair_count"], 0)

    def test_stale_source_and_edited_fingerprint_fail_closed(self):
        path = self.deck()
        fp = self.fp(path)
        tampered = copy.deepcopy(fp)
        tampered["slides"][0]["shapes"][0]["box"][0] = 0.9
        tampered["fingerprint_sha256"] = audit.data_digest(tampered, "fingerprint_sha256")
        with self.assertRaisesRegex(ValueError, "STALE_FINGERPRINT"):
            audit.verify_fingerprint(tampered)
        with ZipFile(path, "a") as archive:
            archive.writestr("changed.txt", "revision")
        with self.assertRaisesRegex(ValueError, "STALE_SOURCE"):
            audit.verify_fingerprint(fp)

    def test_report_verification_rechecks_actual_source_and_comparison(self):
        old, new = self.deck("old.pptx"), self.deck("new.pptx")
        report = self.compare(new, [old])
        self.assertEqual(audit.verify_report(report)["status"], "REJECT_IDENTICAL")
        report["status"] = "MAIN_REVIEW_REQUIRED"
        report["report_sha256"] = audit.data_digest(report, "report_sha256")
        with self.assertRaisesRegex(ValueError, "stale comparison"):
            audit.verify_report(report)
        with ZipFile(old, "a") as archive:
            archive.writestr("changed.txt", "revision")
        with self.assertRaisesRegex(ValueError, "STALE_SOURCE"):
            audit.verify_report(report)

    def test_explicit_output_is_exclusive_and_never_overwrites_source(self):
        source = self.deck()
        before = audit.digest(source)
        output = self.base / "fingerprint.json"
        code, value = self.cli("extract", source, "--deck-id", "current", "--output", output)
        self.assertEqual((code, value["status"]), (0, "EXTRACTED"))
        saved = output.read_bytes()
        code, value = self.cli("extract", source, "--deck-id", "current", "--output", output)
        self.assertEqual((code, value["status"]), (1, "BLOCKED"))
        self.assertEqual(output.read_bytes(), saved)
        code, value = self.cli("extract", source, "--deck-id", "current", "--output", output, "--overwrite")
        self.assertEqual(code, 0)
        code, value = self.cli("extract", source, "--deck-id", "current", "--output", source, "--overwrite")
        self.assertEqual((code, value["status"]), (1, "BLOCKED"))
        self.assertEqual(audit.digest(source), before)

    def test_malformed_or_missing_inputs_return_json_blocked_without_tracebacks(self):
        invalid = self.base / "invalid.json"
        for content in ("", "{", "null", "[]", "{}"):
            invalid.write_text(content, encoding="utf-8")
            code, value = self.cli("compare", invalid)
            self.assertEqual((code, value["status"]), (1, "BLOCKED"))
        code, value = self.cli("extract", self.base / "missing.pptx", "--deck-id", "deck")
        self.assertEqual((code, value["status"]), (1, "BLOCKED"))
        malformed_pptx = self.base / "malformed.pptx"
        with ZipFile(malformed_pptx, "w") as archive:
            archive.writestr("ppt/presentation.xml", "<malformed")
        code, value = self.cli("extract", malformed_pptx, "--deck-id", "deck")
        self.assertEqual((code, value["status"]), (1, "BLOCKED"))
        self.assertIn("Invalid OOXML", " ".join(value["errors"]))


if __name__ == "__main__":
    unittest.main(verbosity=2)
